from datetime import timedelta
from io import BytesIO
import pytest
from sqlalchemy import select
from openpyxl import load_workbook
from app.models import AuthSession,Record,User,Consent,now
from conftest import login

def record(**extra):
    return {'id':'PR-new','name':'Nouveau producteur','project':'Projet p1','village':'Boussé','year':'2026','status':'Actif','sex':'Femme','activity':'Maraîchage','surface':2.5,'production':1000,'visit':'2026-10-02',**extra}

def test_login_csrf_logout(client):
    assert client.get('/impact/producteurs').status_code==401
    assert client.post('/auth/login',json={'email':'super_admin@example.org','password':'incorrect'}).status_code==401
    response=login(client)
    assert 'HttpOnly' in response.headers['set-cookie']
    assert response.json()['user']['role']=='SUPER_ADMIN'
    assert client.get('/auth/me').status_code==200
    csrf=client.headers.pop('X-CSRF-Token')
    assert client.post('/impact/producteurs',json=record()).status_code==403
    client.headers['X-CSRF-Token']=csrf
    assert client.post('/auth/logout').status_code==200
    assert client.get('/auth/me').status_code==401

def test_hostile_origin_and_expiration(client,database):
    response=client.post('/auth/login',headers={'Origin':'https://hostile.example'},json={'email':'super_admin@example.org','password':'MotDePasse-Test-2026'})
    assert response.status_code==403
    login(client)
    with database() as db:
        session=db.scalar(select(AuthSession));session.expires_at=now()-timedelta(seconds=1);db.commit()
    assert client.get('/auth/me').status_code==401

def test_crud_persistence_validation_and_related_records(client,database):
    login(client)
    assert client.post('/impact/producteurs',json=record(surface=-1)).status_code==422
    assert client.post('/impact/producteurs',json=record()).status_code==201
    assert client.post('/impact/producteurs',json=record()).status_code==409
    assert client.patch('/impact/producteurs/PR-new',json={'surface':3.5}).status_code==200
    with database() as db:assert db.scalar(select(Record).where(Record.external_id=='PR-new')).payload['surface']==3.5
    assert client.get('/impact/producteurs/PR-new').json()['surface']==3.5
    assert client.delete('/impact/producteurs/PR-new').status_code==200
    assert client.get('/impact/producteurs/PR-new').status_code==404
    plot={'id':'PA-new','name':'Parcelle','project':'Projet p1','village':'Boussé','year':'2026','status':'Actif','producerId':'PR-0','surface':1,'crop':'Tomate'}
    assert client.post('/impact/parcelles',json=plot).status_code==201
    assert client.delete('/impact/producteurs/PR-0').status_code==409

@pytest.mark.parametrize('role',['SUPER_ADMIN','DIRECTION','CHEF_PROJET','AGENT_TERRAIN','SUIVI_EVALUATION','COMMUNICATION','PARTENAIRE'])
def test_roles_and_project_scope(client,role):
    response=login(client,role)
    assert response.json()['user']['role']==role
    response=client.get('/impact/producteurs')
    if role in ['COMMUNICATION','PARTENAIRE']:
        assert response.status_code==403
    else:
        assert response.status_code==200
        assert len(response.json())==(3 if role in ['SUPER_ADMIN','DIRECTION','SUIVI_EVALUATION'] else 2)
    assert client.get('/impact/utilisateurs').status_code==(200 if role=='SUPER_ADMIN' else 403)
    if role in ['CHEF_PROJET','AGENT_TERRAIN']:
        assert client.get('/impact/producteurs/PR-2').status_code==404
        assert client.post('/impact/producteurs',json=record(project='Projet p2')).status_code==403

def test_server_reports_pdf_excel_and_filters(client):
    login(client)
    filters={'project':'Projet p1','period':'2026','zone':'Toutes les zones','indicators':['Producteurs','Superficie','Femmes','Production']}
    result=client.post('/reports',json=filters)
    assert result.status_code==200,result.text
    assert [r['value'] for r in result.json()['rows']]==[2,3,50,300]
    pdf=client.post('/reports/export/pdf',json=filters);assert pdf.status_code==200;assert pdf.content.startswith(b'%PDF')
    excel=client.post('/reports/export/xlsx',json=filters);assert excel.status_code==200
    workbook=load_workbook(BytesIO(excel.content));assert workbook['Indicateurs']['B2'].value==2;assert workbook['Indicateurs']['B3'].value==3
    zone=client.post('/reports',json={**filters,'zone':'Toéghin'}).json();assert zone['rows'][0]['value']==1
    empty=client.post('/reports',json={**filters,'period':'2024'}).json();assert empty['rows'][0]['value']==0
    login(client,'PARTENAIRE')
    assert client.post('/reports',json=filters).status_code==200
    assert client.post('/reports',json={**filters,'project':'Projet p2'}).status_code==403
    assert client.get('/impact/dashboard').json()['visits']==[]

def test_consent_tracked_and_revoked(client,database):
    login(client)
    assert client.get('/impact/producteurs/PR-0/economic.pdf').status_code==403
    assert client.post('/impact/producteurs/PR-0/consent',json={'accepted':True,'purpose':'Préparation du dossier économique'}).status_code==200
    assert client.get('/impact/producteurs/PR-0/economic.pdf').content.startswith(b'%PDF')
    with database() as db:assert db.scalar(select(Consent)).accepted
    assert client.post('/impact/producteurs/PR-0/consent',json={'accepted':False,'purpose':'Retrait du consentement'}).status_code==200
    assert client.get('/impact/producteurs/PR-0/economic.pdf').status_code==403

def test_user_accounts_and_last_admin(client):
    login(client)
    payload={'id':'agent-new','name':'Agent créé','email':'new@example.org','role':'AGENT_TERRAIN','password':'MotDePasse-Nouveau-2026','projectIds':['p1']}
    response=client.post('/impact/utilisateurs',json=payload);assert response.status_code==201,response.text
    assert 'password' not in response.json() and 'password_hash' not in response.json()
    assert client.patch('/impact/utilisateurs/SUPER_ADMIN',json={'name':'Admin','email':'super_admin@example.org','role':'PARTENAIRE','status':'Actif'}).status_code==409
    assert client.delete('/impact/utilisateurs/SUPER_ADMIN').status_code==409
    assert client.delete('/impact/utilisateurs/agent-new').status_code==200
    assert client.post('/auth/login',json={'email':'new@example.org','password':'MotDePasse-Nouveau-2026'}).status_code==401

def test_public_privacy_contacts_and_newsletter(client):
    catalog=client.get('/public/catalog').json()
    assert 'latitude' not in str(catalog) and 'email' not in str(catalog) and 'PR-0' not in str(catalog)
    assert client.post('/contact',json={'name':'Visiteur','email':'visitor@example.org','message':'Demande de partenariat'}).status_code==201
    assert client.post('/newsletter',json={'email':'visitor@example.org'}).status_code==200
    assert client.get('/impact/messages').status_code==401
    login(client)
    assert len(client.get('/impact/messages').json())==1
    assert len(client.get('/impact/newsletter').json())==1

def test_password_change_revokes_sessions(client):
    login(client)
    assert client.post('/auth/change-password',json={'current_password':'MotDePasse-Test-2026','new_password':'Nouveau-MotDePasse-2026'}).status_code==200
    assert client.get('/auth/me').status_code==401
    assert client.post('/auth/login',json={'email':'super_admin@example.org','password':'MotDePasse-Test-2026'}).status_code==401
    assert client.post('/auth/login',json={'email':'super_admin@example.org','password':'Nouveau-MotDePasse-2026'}).status_code==200

def test_login_rate_limit(client):
    for _ in range(10):assert client.post('/auth/login',json={'email':'missing@example.org','password':'incorrect'}).status_code==401
    assert client.post('/auth/login',json={'email':'missing@example.org','password':'incorrect'}).status_code==429

def test_json_file_transport_preserves_auth_and_consent(client):
    import base64
    filters={'project':'Projet p1','period':'2026','zone':'Toutes les zones','indicators':['Producteurs']}
    assert client.post('/reports/export/pdf?transport=json',json=filters).status_code==401
    login(client)
    response=client.post('/reports/export/pdf?transport=json',json=filters)
    assert response.status_code==200
    assert response.json()['mediaType']=='application/pdf'
    assert base64.b64decode(response.json()['content']).startswith(b'%PDF')
    excel=client.post('/reports/export/xlsx?transport=json',json=filters).json()
    assert load_workbook(BytesIO(base64.b64decode(excel['content'])))['Indicateurs']['B2'].value==2
    url='/impact/producteurs/PR-0/economic.pdf?transport=json'
    assert client.get(url).status_code==403
    client.post('/impact/producteurs/PR-0/consent',json={'accepted':True,'purpose':'Export'})
    assert base64.b64decode(client.get(url).json()['content']).startswith(b'%PDF')
    client.post('/impact/producteurs/PR-0/consent',json={'accepted':False,'purpose':'Retrait'})
    assert client.get(url).status_code==403

def test_private_map_project_scope_and_personal_data_roles(client):
    assert client.get('/impact/map?layer=Producteurs').status_code==401
    login(client,'AGENT_TERRAIN')
    markers=client.get('/impact/map?layer=Producteurs').json()
    assert len(markers)==2 and all(m['project']=='p1' for m in markers)
    assert client.get('/impact/map?layer=Producteurs&project=p2').json()==[]
    login(client,'PARTENAIRE')
    assert client.get('/impact/map?layer=Producteurs').status_code==403
    login(client,'COMMUNICATION')
    assert client.get('/impact/map?layer=Producteurs').status_code==403
