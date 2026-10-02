from io import BytesIO
from datetime import datetime, timezone
from html import escape
from openpyxl import Workbook
from openpyxl.styles import PatternFill, Font
from reportlab.lib import colors
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.pagesizes import A4
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
from fastapi import HTTPException
from sqlalchemy import select
from sqlalchemy.orm import Session
from .models import Project, Record, User
from .schemas import ReportIn
from .security import check_project
from .config import settings

def calculate(db: Session,user: User,filters: ReportIn):
    project=db.scalar(select(Project).where(Project.name==filters.project))
    if not project:raise HTTPException(404,'Projet introuvable.')
    check_project(user,project.id)
    if user.role=='COMMUNICATION':raise HTTPException(403,'Rapports internes non autorisés.')
    statement=select(Record).where(Record.project_id==project.id,Record.year==filters.period,Record.status!='Archivé')
    if filters.zone!='Toutes les zones':statement=statement.where(Record.village==filters.zone)
    rows=db.scalars(statement).all()
    producers=[r.payload for r in rows if r.collection=='producteurs']
    total=lambda field:sum(float(r.get(field,0) or 0) for r in producers)
    values={'Producteurs':(len(producers),'personnes'),'Ménages':(len([r for r in rows if r.collection=='menages']),'ménages'),'Superficie':(round(total('surface'),2),'ha'),'Femmes':(round(100*len([r for r in producers if r.get('sex')=='Femme'])/len(producers),1) if producers else 0,'%'),'Production':(total('production'),'kg'),'Transformateurs':(len([r for r in rows if r.collection=='transformateurs']),'personnes')}
    return {'title':'Rapport d’impact','project':project.name,'period':filters.period,'zone':filters.zone,'source':'Données de travail non validées — résultats non officiels' if settings.seed_work_data else 'Données enregistrées en base — validation institutionnelle requise','generatedAt':datetime.now(timezone.utc).isoformat(),'rows':[{'indicator':name,'value':values[name][0],'unit':values[name][1]} for name in filters.indicators]}

def pdf_bytes(report: dict):
    output=BytesIO();styles=getSampleStyleSheet()
    elements=[Paragraph('CECA-DR — '+escape(report['title']),styles['Title']),Spacer(1,15),Paragraph(escape(report['project']),styles['Heading2']),Paragraph(escape(f"Période : {report['period']} · Zone : {report['zone']}"),styles['Normal']),Spacer(1,10),Paragraph(escape(report['source']),styles['Normal']),Spacer(1,20)]
    table=Table([['Indicateur','Valeur','Unité']]+[[r['indicator'],str(r['value']),r['unit']] for r in report['rows']],colWidths=[250,90,90],repeatRows=1)
    table.setStyle(TableStyle([('BACKGROUND',(0,0),(-1,0),colors.HexColor('#0B5D3B')),('TEXTCOLOR',(0,0),(-1,0),colors.white),('GRID',(0,0),(-1,-1),.3,colors.HexColor('#DCE7DF')),('PADDING',(0,0),(-1,-1),8)]))
    elements.extend([table,Spacer(1,20),Paragraph(escape('Généré : '+report['generatedAt']),styles['Normal'])])
    SimpleDocTemplate(output,pagesize=A4,title=report['title'],author='CECA-DR').build(elements)
    return output.getvalue()

def excel_bytes(report: dict):
    workbook=Workbook();context=workbook.active;context.title='Contexte'
    context.append(['Champ','Valeur'])
    for key in ['project','period','zone','source','generatedAt']:
        value=str(report[key]);context.append([key,"'"+value if value.startswith(('=','+','-','@')) else value])
    sheet=workbook.create_sheet('Indicateurs');sheet.append(['Indicateur','Valeur','Unité'])
    for row in report['rows']:sheet.append([row['indicator'],row['value'],row['unit']])
    for worksheet in [context,sheet]:
        worksheet.freeze_panes='A2'
        worksheet.column_dimensions['A'].width=35;worksheet.column_dimensions['B'].width=45;worksheet.column_dimensions['C'].width=20
        for cell in worksheet[1]:cell.fill=PatternFill('solid',fgColor='0B5D3B');cell.font=Font(bold=True,color='FFFFFF')
    output=BytesIO();workbook.save(output);return output.getvalue()
